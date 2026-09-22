#!/usr/bin/env python3
"""Export the fiscal-publication manifest as a public review workbook.

The source manifest is produced from UOPKOFDA01/UOPKOFDA02 publication pages.
Rows are publication records selected by the local downloader's rule, not every
attachment shown by the site when table-of-contents files were intentionally
skipped.
"""
from __future__ import annotations

from collections import Counter, defaultdict
import argparse
import csv
from datetime import datetime, timezone
from pathlib import Path
import re
import sys
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "openfiscal_publications/lists/manifest.csv"
OUTPUT = ROOT / "docs/data/publication_full_list.xlsx"
SOURCE_URL = "https://www.openfiscaldata.go.kr/op/ko/fd/UOPKOFDA01"
DETAIL_URL = "https://www.openfiscaldata.go.kr/op/ko/fd/UOPKOFDA02"

LIST_SHEET = "재정간행물 전체 목록"
CODE_SHEET = "유형코드표"
NOTE_SHEET = "작성 기준"

LIST_COLUMNS = [
    "연도",
    "유형명",
    "유형코드",
    "간행물명",
    "게시일",
    "게시물ID",
    "파일명",
    "확장자",
    "비고",
    "첨부파일ID",
    "첨부순번",
    "원문URL",
    "선택규칙",
    "상태",
]
CODE_COLUMNS = ["유형코드", "유형명", "record_count", "file_count", "no_attachment_count"]
DANGEROUS_PREFIXES = ("=", "+", "-", "@")


def read_manifest(path: Path = MANIFEST) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def safe_text(value: object) -> str:
    """Return a value safe to write as an Excel text cell.

    Openpyxl will treat strings beginning with '=' as formulas.  Prefixing all
    formula-like values with an apostrophe keeps the value visible as text and
    prevents formula injection if future source data contains such strings.
    """
    if value is None:
        return ""
    text = str(value)
    if text.startswith(DANGEROUS_PREFIXES):
        return "'" + text
    return text


def year_from_date(publication_date: str) -> str:
    match = re.match(r"^(\d{4})", publication_date or "")
    return match.group(1) if match else ""


def filename_for(row: dict[str, str]) -> str:
    for key in ("source_original_name", "source_display_name"):
        value = (row.get(key) or "").strip()
        if value:
            return value
    local_path = (row.get("local_path") or "").strip()
    if local_path:
        return Path(local_path).name
    return ""


def extension_for(filename: str, local_path: str = "") -> str:
    candidate = filename or Path(local_path or "").name
    suffix = Path(candidate).suffix.lower().lstrip(".")
    return suffix


def note_for(row: dict[str, str]) -> str:
    notes: list[str] = []
    status = row.get("status", "")
    if status == "no_downloadable_file":
        notes.append("첨부 없음")
    elif status and status != "downloaded":
        notes.append(status)
    if row.get("error"):
        notes.append(row["error"])
    return "; ".join(notes)


def build_list_rows(rows: Iterable[dict[str, str]]) -> list[list[str]]:
    output: list[list[str]] = []
    for row in rows:
        filename = filename_for(row)
        output.append([
            year_from_date(row.get("publication_date", "")),
            row.get("category_name", ""),
            row.get("category_code", ""),
            row.get("publication_title", ""),
            row.get("publication_date", ""),
            row.get("publication_seq", ""),
            filename,
            extension_for(filename, row.get("local_path", "")),
            note_for(row),
            row.get("atch_file_id", ""),
            row.get("atch_file_seq", ""),
            row.get("detail_url", ""),
            row.get("selection_rule", ""),
            row.get("status", ""),
        ])
    return output


def build_code_rows(rows: Iterable[dict[str, str]]) -> list[list[object]]:
    grouped: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[(row.get("category_code", ""), row.get("category_name", ""))].append(row)

    code_rows: list[list[object]] = []
    for (code, name), members in sorted(grouped.items(), key=lambda item: (int(item[0][0] or 0), item[0][1])):
        file_count = sum(1 for row in members if row.get("status") == "downloaded" and row.get("atch_file_id"))
        no_attachment_count = sum(1 for row in members if row.get("status") == "no_downloadable_file" or not row.get("atch_file_id"))
        code_rows.append([code, name, len(members), file_count, no_attachment_count])
    return code_rows


def build_note_rows(rows: list[dict[str, str]]) -> list[tuple[str, str]]:
    statuses = Counter(row.get("status", "") for row in rows)
    return [
        ("작성 시각(UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
        ("원천 목록", str(MANIFEST.relative_to(ROOT))),
        ("원천 화면", SOURCE_URL),
        ("상세 화면", DETAIL_URL),
        ("데이터 성격", "UOPKOFDA01 재정간행물의 기존 수집 스냅샷 전체 레코드 목록. 엑셀 작성일에 사이트 전체를 실시간 재수집한 결과가 아님"),
        ("전체 레코드", str(len(rows))),
        ("다운로드 성공 파일", str(statuses.get("downloaded", 0))),
        ("첨부 없음 레코드", str(statuses.get("no_downloadable_file", 0))),
        ("유형 수", str(len({row.get("category_code", "") for row in rows}))),
        ("유형코드 출처", "서버 분류코드 OP061(상위 OP060=05)의 itgDtsCd이며 상세 URL 파라미터 ofdBrdiDtsClsCd로도 제공됨"),
        ("연도 기준", "게시일(publication_date)의 앞 4자리. 제목의 연도는 추정에 사용하지 않음"),
        ("파일명 기준", "manifest.csv의 source_original_name 우선, 없으면 source_display_name/local_path 순서"),
        ("선택 기준", "원문이 있으면 원문 우선·목차 제외. 원문이 없으면 표지 이외의 제공 파일(목차 분할본 포함)을 선택. 따라서 사이트의 모든 표지·첨부를 의미하지 않음"),
        ("ZIP 기준", "ZIP 내부 파일은 별도 행으로 펼치지 않고 ZIP 첨부 1건으로 계산"),
        ("검토 제외 기준", "type-examples 페이지의 브라우저 localStorage 제외 선택은 이 엑셀에 반영하지 않음"),
        ("수식 안전", "원천 문자열이 =,+,-,@ 로 시작하면 텍스트로 보이도록 앞에 apostrophe를 붙임"),
    ]


def style_table(ws, freeze: str = "A2") -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 80))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[col_letter].width = max(10, min(max_len + 2, 60))


def write_workbook(rows: list[dict[str, str]], output: Path = OUTPUT) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = LIST_SHEET
    ws.append(LIST_COLUMNS)
    for record in build_list_rows(rows):
        ws.append([safe_text(value) for value in record])
    style_table(ws)

    code_ws = wb.create_sheet(CODE_SHEET)
    code_ws.append(CODE_COLUMNS)
    for record in build_code_rows(rows):
        code_ws.append([safe_text(record[0]), safe_text(record[1]), *record[2:]])
    style_table(code_ws)

    note_ws = wb.create_sheet(NOTE_SHEET)
    note_ws.append(["항목", "내용"])
    for key, value in build_note_rows(rows):
        note_ws.append([safe_text(key), safe_text(value)])
    style_table(note_ws)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    wb.save(temporary)
    temporary.replace(output)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=MANIFEST)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> None:
    args = parse_args(list(sys.argv[1:] if argv is None else argv))
    rows = read_manifest(args.manifest)
    write_workbook(rows, args.output)
    statuses = Counter(row.get("status", "") for row in rows)
    print(
        f"wrote {args.output} records={len(rows)} "
        f"downloaded={statuses.get('downloaded', 0)} "
        f"no_downloadable_file={statuses.get('no_downloadable_file', 0)} "
        f"types={len({row.get('category_code', '') for row in rows})}"
    )


if __name__ == "__main__":
    main()
