from __future__ import annotations

import hashlib
import tempfile
from pathlib import Path
import unittest

from openpyxl import load_workbook

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import export_publication_list as subject


class PublicationListExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.rows = subject.read_manifest()

    def test_manifest_scope_counts_and_codes(self) -> None:
        self.assertEqual(len(self.rows), 1698)
        self.assertEqual(len({row["category_code"] for row in self.rows}), 15)
        self.assertEqual(sum(1 for row in self.rows if row["status"] == "downloaded"), 1695)
        self.assertEqual(sum(1 for row in self.rows if row["status"] == "no_downloadable_file"), 3)
        self.assertEqual({row["publication_id"] for row in self.rows}, {"FSL1004"})
        self.assertIn("501", {row["category_code"] for row in self.rows})
        self.assertIn("519", {row["category_code"] for row in self.rows})

    def test_code_summary_distinguishes_records_files_and_missing(self) -> None:
        summary = {row[0]: row for row in subject.build_code_rows(self.rows)}
        self.assertEqual(len(summary), 15)
        self.assertEqual(sum(row[2] for row in summary.values()), 1698)
        self.assertEqual(sum(row[3] for row in summary.values()), 1695)
        self.assertEqual(sum(row[4] for row in summary.values()), 3)
        self.assertEqual(summary["505"][4], 2)
        self.assertEqual(summary["516"][4], 1)

    def test_list_rows_are_complete_and_use_source_original_filename(self) -> None:
        rows = subject.build_list_rows(self.rows)
        self.assertEqual(len(rows), 1698)
        self.assertEqual(len(rows[0]), len(subject.LIST_COLUMNS))
        first = rows[0]
        by_header = dict(zip(subject.LIST_COLUMNS, first))
        self.assertEqual(by_header["연도"], "2026")
        self.assertEqual(by_header["유형코드"], "501")
        self.assertEqual(by_header["파일명"], "기획예산처 월간 재정동향 26.6월호..pdf")
        self.assertEqual(by_header["확장자"], "pdf")
        no_attachment = [dict(zip(subject.LIST_COLUMNS, row)) for row in rows if row[13] == "no_downloadable_file"]
        self.assertEqual(len(no_attachment), 3)
        self.assertTrue(all(row["비고"].startswith("첨부 없음") for row in no_attachment))
        self.assertTrue(all(row["파일명"] == "" for row in no_attachment))

    def test_workbook_layout_counts_and_notes(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "publication_full_list.xlsx"
            subject.write_workbook(self.rows, output)
            wb = load_workbook(output, data_only=False)
            self.assertEqual(wb.sheetnames, [subject.LIST_SHEET, subject.CODE_SHEET, subject.NOTE_SHEET])
            self.assertEqual(wb[subject.LIST_SHEET].max_row, 1699)
            self.assertEqual(wb[subject.CODE_SHEET].max_row, 16)
            self.assertEqual(wb[subject.LIST_SHEET]["A2"].data_type, "s")
            notes = {row[0].value: row[1].value for row in wb[subject.NOTE_SHEET].iter_rows(min_row=2, max_col=2)}
            self.assertEqual(notes["전체 레코드"], "1698")
            self.assertEqual(notes["다운로드 성공 파일"], "1695")
            self.assertIn("원문이 있으면 원문 우선·목차 제외", notes["선택 기준"])
            self.assertIn("목차 분할본 포함", notes["선택 기준"])
            self.assertIn("실시간 재수집한 결과가 아님", notes["데이터 성격"])
            self.assertIn("OP061", notes["유형코드 출처"])

    def test_safe_text_blocks_formula_like_source_strings(self) -> None:
        for raw in ("=SUM(1,1)", "+1", "-1", "@cmd"):
            self.assertTrue(subject.safe_text(raw).startswith("'"))
        self.assertEqual(subject.safe_text("정상값"), "정상값")

    def test_export_does_not_mutate_source_manifest(self) -> None:
        before = hashlib.sha256(subject.MANIFEST.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as tmpdir:
            subject.main(["--output", str(Path(tmpdir) / "publication_full_list.xlsx")])
        after = hashlib.sha256(subject.MANIFEST.read_bytes()).hexdigest()
        self.assertEqual(after, before)


if __name__ == "__main__":
    unittest.main()
